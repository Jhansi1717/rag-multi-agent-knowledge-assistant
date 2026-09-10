"""ingestion/docx_extractor.py

Phase 3.3 — DOCX Extraction via python-docx

Provides :class:`DocxExtractor` — the focused implementation of
DOCX-to-text extraction for Milestone 1.

Design decisions
----------------
* Uses ``python-docx`` (``docx``).
* Iterates elements in document body order (preserving the relative order
  between paragraphs and tables).
* Non-empty paragraphs and tables are extracted as textual blocks.
* Tables are formatted as clean, human-readable pipe-separated rows.
* Records metadata including ``paragraph_count``, ``table_count``,
  ``file_size_bytes``, and ``has_text``.
* Returns a dictionary with ``text`` and ``metadata`` matching the common ingestion interface.
"""

from __future__ import annotations

import os
from typing import Any, Dict, List

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph


class DocxExtractor:
    """Extract text and metadata from a DOCX file using python-docx.

    Usage
    -----
    ::

        extractor = DocxExtractor()
        result = extractor.extract("data/technical/api_documentation.docx")
        print(result["text"])
        print(result["metadata"]["paragraph_count"])
    """

    def extract(self, file_path: str) -> Dict[str, Any]:
        """Open *file_path*, extract paragraphs and tables, return a result dict.

        Parameters
        ----------
        file_path:
            Path to a ``.docx`` file. Must exist and be a regular file.

        Returns
        -------
        dict with keys:

        ``text`` : str
            Full text of the document (paragraphs/tables joined by ``"\\n\\n"``).

        ``metadata`` : dict with keys:
            ``source_location`` : str
            ``document_name`` : str
            ``file_size_bytes`` : int
            ``format`` : str (always "docx")
            ``paragraph_count`` : int
            ``table_count`` : int
            ``has_text`` : bool

        Raises
        ------
        FileNotFoundError
            If *file_path* does not exist.
        RuntimeError
            If python-docx cannot open the file.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path!r}")

        filename = os.path.basename(file_path)
        file_size = os.path.getsize(file_path)

        try:
            doc = docx.Document(file_path)
        except Exception as exc:
            raise RuntimeError(
                f"python-docx could not open {file_path!r}: {exc}"
            ) from exc

        content_blocks: List[str] = []
        paragraph_count = len(doc.paragraphs)
        table_count = len(doc.tables)

        # Preserve document body order by iterating child XML elements
        for child in doc.element.body:
            tag = child.tag
            if tag.endswith("p"):
                para = Paragraph(child, doc)
                text = para.text.strip()
                if text:
                    content_blocks.append(text)
            elif tag.endswith("tbl"):
                tbl = Table(child, doc)
                table_lines: List[str] = []
                for row in tbl.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(row_cells):
                        table_lines.append(" | ".join(row_cells))
                if table_lines:
                    content_blocks.append("\n".join(table_lines))

        full_text = "\n\n".join(content_blocks)
        has_text = len(full_text.strip()) > 0

        metadata: Dict[str, Any] = {
            "source_location": file_path,
            "document_name": filename,
            "file_size_bytes": file_size,
            "format": "docx",
            "paragraph_count": paragraph_count,
            "table_count": table_count,
            "has_text": has_text,
        }

        return {
            "text": full_text,
            "metadata": metadata,
        }
