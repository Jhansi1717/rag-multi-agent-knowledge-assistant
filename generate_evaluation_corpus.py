"""Generate M1 evaluation corpus — Software Engineering & Hospital Administration."""

import os

import pandas as pd
from docx import Document
from fpdf import FPDF

SOFTWARE_DIR = "data/software_engineering"
HOSPITAL_DIR = "data/hospital_administration"
EVAL_DIR = "data/evaluation"


def _write_pdf(path: str, title: str, sections: list[tuple[str, str]]) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=14)
    pdf.cell(0, 10, title, ln=1, align="C")
    pdf.ln(4)
    pdf.set_font("Arial", size=11)
    for heading, body in sections:
        pdf.set_font("Arial", style="B", size=11)
        pdf.multi_cell(0, 7, heading)
        pdf.set_font("Arial", size=11)
        pdf.multi_cell(0, 6, body)
        pdf.ln(2)
    pdf.output(path)


def generate_software_engineering() -> None:
    os.makedirs(SOFTWARE_DIR, exist_ok=True)

    _write_pdf(
        f"{SOFTWARE_DIR}/microservices_architecture.pdf",
        "Microservices and Agile Delivery Guide",
        [
            (
                "1. Agile Sprint Cadence",
                "Teams operate in two-week sprints. Each sprint begins with planning, "
                "includes daily standups, and ends with a retrospective. Sprint goals "
                "must be demo-ready by the review meeting.",
            ),
            (
                "2. Microservices vs Monolith",
                "Microservices decompose applications into independently deployable services "
                "with bounded contexts. Monolithic architectures keep all modules in one "
                "deployable unit. Microservices improve scalability but add network complexity.",
            ),
            (
                "3. Continuous Integration",
                "Every pull request triggers automated unit tests, lint checks, and security "
                "scans. Builds must pass before merge to the main branch.",
            ),
            (
                "4. Service Communication",
                "Services communicate via REST APIs or asynchronous message queues. "
                "Synchronous calls require circuit breakers and timeout policies.",
            ),
        ],
    )

    doc = Document()
    doc.add_heading("Engineering Coding Standards", 0)
    doc.add_heading("Naming Conventions", level=1)
    doc.add_paragraph(
        "Classes use PascalCase. Functions and variables use snake_case. "
        "Constants use UPPER_SNAKE_CASE. Module names are lowercase without underscores."
    )
    doc.add_heading("Error Handling", level=1)
    doc.add_paragraph(
        "Catch exceptions at service boundaries. Log errors with structured JSON including "
        "timestamp, severity, correlation_id, and stack trace. Never swallow exceptions silently."
    )
    doc.add_heading("Documentation", level=1)
    doc.add_paragraph(
        "Public APIs require docstrings describing parameters, return values, and raised exceptions. "
        "README files must include setup, test, and deployment instructions."
    )
    doc.add_heading("Code Review Checklist", level=1)
    table = doc.add_table(rows=4, cols=2)
    rows = [
        ("Readability", "Code is self-explanatory with minimal comments"),
        ("Tests", "New logic includes unit tests"),
        ("Security", "No hardcoded secrets or credentials"),
        ("Performance", "No obvious N+1 queries or blocking calls"),
    ]
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
    doc.save(f"{SOFTWARE_DIR}/coding_standards.docx")

    git_content = """Git Workflow for Software Teams

1. Branch Naming
Feature branches use the format feature/<ticket-id>-short-description.
Bug fixes use bugfix/<ticket-id>-short-description.
Hotfixes branch from main using hotfix/<ticket-id>.

2. Creating a Feature Branch
Step 1: Pull latest changes from the develop branch.
Step 2: Create a new branch: git checkout -b feature/123-add-login-form develop
Step 3: Commit changes with conventional commit messages.
Step 4: Push the branch and open a pull request.

3. Pull Request Rules
Every pull request requires at least two approving reviews.
All CI checks must pass before merge.
Squash merge is preferred for feature branches.

4. Release Process
Releases are tagged on main using semantic versioning (vMAJOR.MINOR.PATCH).
Release notes are generated from merged pull request titles.
"""
    with open(f"{SOFTWARE_DIR}/git_workflow.txt", "w", encoding="utf-8") as f:
        f.write(git_content)

    pd.DataFrame(
        [
            ["Python", "Multi-paradigm", "Dynamic", "Data science and backends", "Low"],
            ["Java", "Object-oriented", "Static", "Enterprise applications", "Medium"],
            ["JavaScript", "Multi-paradigm", "Dynamic", "Web frontends", "Low"],
            ["Rust", "Systems", "Static", "Systems programming", "High"],
            ["Go", "Procedural", "Static", "Cloud infrastructure", "Medium"],
        ],
        columns=["Language", "Paradigm", "Typing", "Primary_Use_Case", "Learning_Curve"],
    ).to_csv(f"{SOFTWARE_DIR}/programming_languages.csv", index=False)


def generate_hospital_administration() -> None:
    os.makedirs(HOSPITAL_DIR, exist_ok=True)

    _write_pdf(
        f"{HOSPITAL_DIR}/patient_admission_policy.pdf",
        "Riverside General Hospital - Patient Admission Policy",
        [
            (
                "1. Registration Requirements",
                "Patients must present a government-issued photo ID and insurance card at admission. "
                "Emergency admissions may proceed with partial information but must complete "
                "registration within 24 hours.",
            ),
            (
                "2. Insurance Verification",
                "Admitting staff verify coverage through the electronic eligibility system before "
                "assigning an inpatient bed. Pre-authorization is required for elective surgeries.",
            ),
            (
                "3. Admission Paperwork Timeline",
                "Non-emergency patients must finalize admission paperwork within 4 hours of arrival. "
                "Incomplete forms are flagged for patient financial services follow-up.",
            ),
            (
                "4. Discharge Planning",
                "Discharge planning begins within 24 hours of admission for expected stays longer "
                "than 48 hours. Social work consults are available for transitional care.",
            ),
        ],
    )

    doc = Document()
    doc.add_heading("Clinical Nursing Procedures Manual", 0)
    doc.add_heading("Vital Signs Monitoring", level=1)
    doc.add_paragraph(
        "Measure temperature, pulse, respiration, and blood pressure every 4 hours for "
        "standard inpatient units. Critical care units require continuous monitoring."
    )
    doc.add_heading("Medication Administration", level=1)
    doc.add_paragraph(
        "Follow the five rights: right patient, right drug, right dose, right route, right time. "
        "Scan the patient wristband and medication barcode before administration. "
        "Document administration immediately in the electronic health record."
    )
    doc.add_heading("Hand Hygiene Protocol", level=1)
    doc.add_paragraph(
        "Perform hand hygiene using alcohol-based rub for 20 seconds before and after "
        "every patient contact. Use soap and water when hands are visibly soiled."
    )
    doc.add_heading("Patient Charting Standards", level=1)
    doc.add_paragraph(
        "Chart entries must be timely, accurate, and signed with credentials. "
        "Late entries require an explanation note."
    )
    doc.save(f"{HOSPITAL_DIR}/nursing_procedures.docx")

    emergency_content = """Riverside General Hospital - Emergency Response Protocols

1. Code Blue (Cardiac Arrest)
Activate Code Blue for unresponsive patients without a pulse.
Step 1: Call the emergency response team via the overhead alert.
Step 2: Begin CPR at 100-120 compressions per minute.
Step 3: Apply AED pads and follow device prompts.
Step 4: Assign roles for airway, medications, and documentation.

2. Code Red (Fire)
Code Red indicates an active fire. RACE protocol applies:
Rescue patients in immediate danger.
Activate the alarm and call security.
Confine the fire by closing doors.
Extinguish only if safe; otherwise evacuate.

3. Disaster Triage Levels
Triage Level 1 (Red): Immediate life-threatening injuries.
Triage Level 2 (Yellow): Urgent but stable conditions.
Triage Level 3 (Green): Minor injuries, can wait.
Triage Level 4 (Black): Deceased or expectant cases.

4. Mass Casualty Activation
Mass casualty incidents require opening the surge capacity plan and notifying the incident commander.
"""
    with open(f"{HOSPITAL_DIR}/emergency_protocols.txt", "w", encoding="utf-8") as f:
        f.write(emergency_content)

    pd.DataFrame(
        [
            ["Acetaminophen", "650 mg", "Oral", "Every 6 hours", "3000 mg"],
            ["Ibuprofen", "400 mg", "Oral", "Every 8 hours", "1200 mg"],
            ["Ibuprofen", "800 mg", "Intravenous", "Every 8 hours", "2400 mg"],
            ["Ondansetron", "4 mg", "Intravenous", "Every 8 hours", "12 mg"],
            ["Normal Saline", "1000 mL", "Intravenous", "As needed", "3000 mL"],
        ],
        columns=["Medication", "Standard_Dose", "Route", "Frequency", "Max_Daily_Dose"],
    ).to_csv(f"{HOSPITAL_DIR}/medication_dosage_reference.csv", index=False)


def main() -> None:
    os.makedirs(EVAL_DIR, exist_ok=True)
    generate_software_engineering()
    generate_hospital_administration()
    print("Evaluation corpus generated.")


if __name__ == "__main__":
    main()
