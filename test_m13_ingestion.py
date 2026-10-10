"""M1.3 — Focused ingestion pipeline tests (all formats, chunking, metadata, dedupe)."""

import os
import uuid

import pandas as pd
import numpy as np
import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from ingestion import (
    DEFAULT_CHUNK_OVERLAP,
    DEFAULT_CHUNK_SIZE,
    TextChunker,
    ingest_file,
    index_file,
)
from ingestion.models import SUPPORTED_EXTENSIONS, make_chunk, make_document
from vector_store.store import VectorStore


@pytest.fixture
def tmp_data_dir(tmp_path):
    hr = tmp_path / "data" / "hr"
    hr.mkdir(parents=True)
    return hr


@pytest.fixture
def embedder():
    """Shared lightweight embedder for vector store tests."""
    from vector_store.embeddings import EmbeddingProvider
    return EmbeddingProvider()


def _write_pdf(path, lines):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in lines:
        pdf.multi_cell(0, 8, txt=line)
    pdf.output(str(path))


def _write_docx(path, paragraphs, table_rows=None):
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, val in enumerate(row):
                table.rows[r].cells[c].text = str(val)
    doc.save(str(path))


def _write_txt(path, content):
    path.write_text(content, encoding="utf-8")


def _write_csv(path, rows, columns):
    pd.DataFrame(rows, columns=columns).to_csv(path, index=False)


def _assert_chunk_metadata(chunk, filename, file_type, domain):
    meta = chunk.to_dict()
    assert meta["chunk_id"]
    assert meta["document_id"]
    assert meta["chunk_index"] >= 0
    assert meta["filename"] == filename
    assert meta["file_type"] == file_type
    assert meta["domain"] == domain
    assert meta["text"].strip()
    assert chunk.metadata.get("token_count", 0) > 0


class TestFormatExtraction:
    def test_empty_text_document_is_rejected(self, tmp_data_dir):
        path = tmp_data_dir / "empty.txt"
        _write_txt(path, " \n\t\n")

        with pytest.raises(ValueError, match="No extractable text"):
            ingest_file(str(path))

    def test_corrupt_pdf_is_rejected(self, tmp_data_dir):
        path = tmp_data_dir / "corrupt.pdf"
        path.write_bytes(b"not a valid PDF")

        with pytest.raises(RuntimeError, match="could not open"):
            ingest_file(str(path))

    def test_pdf_extraction_and_page_metadata(self, tmp_data_dir):
        path = tmp_data_dir / "policy.pdf"
        _write_pdf(path, ["Page one content about leave.", "Page two content about sick days."])
        doc, chunks = ingest_file(str(path), domain="HR")
        assert doc.file_type == ".pdf"
        assert doc.metadata["page_count"] >= 1
        assert "page_texts" in doc.metadata
        assert len(chunks) >= 1
        assert any(c.metadata.get("page_number") or c.metadata.get("page_start") for c in chunks)

    def test_docx_paragraphs_and_tables(self, tmp_data_dir):
        path = tmp_data_dir / "api.docx"
        _write_docx(
            path,
            ["REST API overview.", "Use DELETE to remove resources."],
            table_rows=[["Method", "Action"], ["DELETE", "Remove"]],
        )
        doc, chunks = ingest_file(str(path), domain="Software")
        assert doc.file_type == ".docx"
        assert doc.metadata["paragraph_count"] >= 2
        assert doc.metadata["table_count"] >= 1
        assert "DELETE" in doc.text
        assert len(chunks) >= 1

    def test_txt_utf8(self, tmp_data_dir):
        path = tmp_data_dir / "notes.txt"
        _write_txt(path, "Line one: café résumé\n\nLine two: naïve coöperate")
        doc, chunks = ingest_file(str(path), domain="HR")
        assert doc.file_type == ".txt"
        assert doc.metadata["line_count"] >= 2
        assert "café" in doc.text
        assert len(chunks) >= 1

    def test_csv_semantic_text(self, tmp_data_dir):
        path = tmp_data_dir / "allowance.csv"
        _write_csv(
            path,
            [["Annual", 20, 5], ["Sick", 10, 0]],
            ["Leave_Type", "Days_Allowed", "Carry_Forward_Limit"],
        )
        doc, chunks = ingest_file(str(path), domain="HR")
        assert doc.file_type == ".csv"
        assert doc.metadata["row_count"] == 2
        assert "Leave_Type: Annual" in doc.text
        assert any(c.metadata.get("row_number") or c.metadata.get("row_start") for c in chunks)


class TestChunking:
    def test_defaults_within_m13_range(self):
        assert 500 <= DEFAULT_CHUNK_SIZE <= 800
        assert 50 <= DEFAULT_CHUNK_OVERLAP <= 100

    def test_paragraph_boundary_preferred(self, tmp_data_dir):
        path = tmp_data_dir / "paragraphs.txt"
        paras = [f"Paragraph {i}. " + ("word " * 120) for i in range(8)]
        _write_txt(path, "\n\n".join(paras))
        _, chunks = ingest_file(str(path))
        assert len(chunks) >= 2
        for chunk in chunks:
            assert "\n\n" in chunk.text or len(chunks) == 1
            assert not chunk.text.startswith(" ")

    def test_chunk_metadata_fields(self, tmp_data_dir):
        path = tmp_data_dir / "meta.pdf"
        _write_pdf(path, ["Employees receive twenty days of annual leave per year."])
        _, chunks = ingest_file(str(path), domain="HR")
        assert chunks
        _assert_chunk_metadata(chunks[0], "meta.pdf", ".pdf", "HR")

    def test_chunking_is_deterministic_and_retains_each_segment(self, tmp_data_dir):
        from ingestion.chunker import chunk_document

        paragraphs = [f"Unique paragraph marker {index}. " + ("detail " * 35) for index in range(12)]
        path = tmp_data_dir / "complete.txt"
        _write_txt(path, "\n\n".join(paragraphs))
        document, _ = ingest_file(str(path), domain="HR")
        first = TextChunker(chunk_size=80, chunk_overlap=15).chunk_document(document)
        second = TextChunker(chunk_size=80, chunk_overlap=15).chunk_document(document)

        assert [(chunk.chunk_id, chunk.text) for chunk in first] == [
            (chunk.chunk_id, chunk.text) for chunk in second
        ]
        for index in range(len(paragraphs)):
            marker = f"Unique paragraph marker {index}."
            assert any(marker in chunk.text for chunk in first)

    def test_repeated_ingestion_keeps_document_and_chunk_ids_stable(self, tmp_data_dir):
        path = tmp_data_dir / "stable.txt"
        _write_txt(path, "Stable identity for repeated ingestion.")

        first_document, first_chunks = ingest_file(str(path), domain="HR")
        second_document, second_chunks = ingest_file(str(path), domain="HR")

        assert second_document.document_id == first_document.document_id
        assert [chunk.chunk_id for chunk in second_chunks] == [
            chunk.chunk_id for chunk in first_chunks
        ]
        assert [chunk.text for chunk in second_chunks] == [chunk.text for chunk in first_chunks]
        assert second_chunks[0].metadata["domain"] == "HR"


class TestVectorIndexing:
    def test_document_identity_replacement_and_reload_preserve_exact_row_mapping(
        self, tmp_data_dir
    ):
        class DeterministicEmbedder:
            model_name = "test-embedder"
            dimension = 2

            def encode_chunks(self, chunks):
                return np.array(
                    [[float(len(chunk.text)), float(chunk.chunk_index)] for chunk in chunks],
                    dtype=np.float32,
                )

            def encode_texts(self, texts):
                return np.array([[float(len(text)), 0.0] for text in texts], dtype=np.float32)

            def encode_query(self, query):
                return np.array([1.0, 0.0], dtype=np.float32)

        index_path = str(tmp_data_dir / "sequence.faiss")
        meta_path = str(tmp_data_dir / "sequence.json")
        store = VectorStore(
            index_path=index_path,
            meta_path=meta_path,
            embedder=DeterministicEmbedder(),
        )

        def make_document_and_chunks(document_id, filename, source, texts, domain):
            document = make_document(
                filename=filename,
                file_type=".txt",
                source=source,
                text="\n\n".join(texts),
                document_id=document_id,
            )
            chunks = [
                make_chunk(
                    document_id=document_id,
                    document_name=filename,
                    chunk_index=index,
                    text=text,
                    source_location=source,
                    metadata={"domain": domain, "row_number": index + 1},
                )
                for index, text in enumerate(texts)
            ]
            return document, chunks

        source_a = str(tmp_data_dir / "a.txt")
        source_b = str(tmp_data_dir / "b.txt")
        document_a, chunks_a = make_document_and_chunks(
            "doc-a", "a.txt", source_a, ["old A one", "old A two"], "Software Engineering"
        )
        document_b, chunks_b = make_document_and_chunks(
            "doc-b", "b.txt", source_b, ["original B"], "Hospital Administration"
        )
        replacement_a, replacement_chunks_a = make_document_and_chunks(
            "doc-a", "a.txt", source_a, ["replacement A"], "Software Engineering"
        )

        assert store.index.ntotal == 0
        store.index_document(document_a, chunks_a)
        store.index_document(document_b, chunks_b)
        assert store.index.ntotal == 3
        store.index_document(replacement_a, replacement_chunks_a)

        expected_records = [
            store.metadata[str(row)] for row in range(store.index.ntotal)
        ]
        assert store.index.ntotal == len(store.metadata) == 2
        assert [record["chunk_id"] for record in expected_records] == [
            "doc-b_0",
            "doc-a_0",
        ]
        for row, record in enumerate(expected_records):
            vector = store.index.reconstruct(row)
            assert vector[0] == float(len(record["text"]))
            assert record["domain"] == record["metadata"]["domain"]

        store.save()
        loaded = VectorStore(
            index_path=index_path,
            meta_path=meta_path,
            embedder=store.embedder,
        )
        loaded.load()
        assert loaded.index.ntotal == len(loaded.metadata) == 2
        assert [loaded.metadata[str(row)] for row in range(2)] == expected_records
        assert [
            loaded.index.reconstruct(row).tolist() for row in range(2)
        ] == [store.index.reconstruct(row).tolist() for row in range(2)]
        assert {hit["document_id"] for hit in loaded.search("query", top_k=2)} == {
            "doc-a",
            "doc-b",
        }

        loaded.index_document(replacement_a, replacement_chunks_a)
        assert loaded.index.ntotal == len(loaded.metadata) == 2
        assert [loaded.metadata[str(row)] for row in range(2)] == expected_records

    def test_replacement_by_document_identity_removes_previous_source(self, tmp_data_dir):
        class DeterministicEmbedder:
            model_name = "test-embedder"
            dimension = 2

            def encode_chunks(self, chunks):
                return np.array([[float(len(chunk.text)), 0.0] for chunk in chunks], dtype=np.float32)

            def encode_texts(self, texts):
                return np.array([[float(len(text)), 0.0] for text in texts], dtype=np.float32)

            def encode_query(self, query):
                return np.array([1.0, 0.0], dtype=np.float32)

        store = VectorStore(
            index_path=str(tmp_data_dir / "identity.faiss"),
            meta_path=str(tmp_data_dir / "identity.json"),
            embedder=DeterministicEmbedder(),
        )
        old_source = str(tmp_data_dir / "old.txt")
        new_source = str(tmp_data_dir / "renamed.txt")
        old_document = make_document(
            filename="old.txt", file_type=".txt", source=old_source,
            text="old identity content", document_id="stable-doc",
        )
        old_chunk = make_chunk(
            document_id="stable-doc", document_name="old.txt", chunk_index=0,
            text="old identity content", source_location=old_source,
        )
        new_document = make_document(
            filename="renamed.txt", file_type=".txt", source=new_source,
            text="new identity content", document_id="stable-doc",
        )
        new_chunk = make_chunk(
            document_id="stable-doc", document_name="renamed.txt", chunk_index=0,
            text="new identity content", source_location=new_source,
        )

        store.index_document(old_document, [old_chunk])
        store.index_document(new_document, [new_chunk])

        assert store.index.ntotal == len(store.metadata) == 1
        assert store.metadata["0"]["source_location"] == new_source
        assert store.metadata["0"]["text"] == "new identity content"

    def test_load_rejects_missing_metadata_row(self, tmp_data_dir):
        import json

        class DeterministicEmbedder:
            model_name = "test-embedder"
            dimension = 2

            def encode_chunks(self, chunks):
                return np.ones((len(chunks), 2), dtype=np.float32)

            def encode_texts(self, texts):
                return np.ones((len(texts), 2), dtype=np.float32)

            def encode_query(self, query):
                return np.ones(2, dtype=np.float32)

        index_path = str(tmp_data_dir / "broken-map.faiss")
        meta_path = tmp_data_dir / "broken-map.json"
        store = VectorStore(
            index_path=index_path,
            meta_path=str(meta_path),
            embedder=DeterministicEmbedder(),
        )
        document = make_document(
            filename="mapped.txt", file_type=".txt", source="mapped.txt",
            text="mapped text", document_id="mapped-doc",
        )
        chunk = make_chunk(
            document_id="mapped-doc", document_name="mapped.txt", chunk_index=0,
            text="mapped text", source_location="mapped.txt",
        )
        store.index_document(document, [chunk])
        store.save()
        payload = json.loads(meta_path.read_text(encoding="utf-8"))
        payload.pop("0")
        meta_path.write_text(json.dumps(payload), encoding="utf-8")

        loaded = VectorStore(
            index_path=index_path,
            meta_path=str(meta_path),
            embedder=store.embedder,
        )
        with pytest.raises(ValueError, match="row IDs do not align"):
            loaded.load()

    def test_failed_metadata_write_restores_persisted_pair(self, tmp_data_dir, monkeypatch):
        import shutil
        import vector_store.store as store_module

        class DeterministicEmbedder:
            model_name = "test-embedder"
            dimension = 2

            def encode_chunks(self, chunks):
                return np.array([[float(len(chunk.text)), 0.0] for chunk in chunks], dtype=np.float32)

            def encode_texts(self, texts):
                return np.array([[float(len(text)), 0.0] for text in texts], dtype=np.float32)

            def encode_query(self, query):
                return np.ones(2, dtype=np.float32)

        index_path = tmp_data_dir / "transaction.faiss"
        meta_path = tmp_data_dir / "transaction.json"
        embedder = DeterministicEmbedder()
        store = VectorStore(
            index_path=str(index_path), meta_path=str(meta_path), embedder=embedder
        )

        def add_document(document_id, filename, text):
            document = make_document(
                filename=filename, file_type=".txt", source=filename,
                text=text, document_id=document_id,
            )
            chunk = make_chunk(
                document_id=document_id, document_name=filename, chunk_index=0,
                text=text, source_location=filename,
            )
            store.index_document(document, [chunk])

        add_document("doc-1", "one.txt", "first persisted document")
        store.save()
        previous_index = index_path.read_bytes()
        previous_metadata = meta_path.read_bytes()
        add_document("doc-2", "two.txt", "second in-memory document")

        original_copyfile = shutil.copyfile
        calls = 0

        def fail_metadata_copy(source, destination, *args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OSError("simulated metadata persistence failure")
            return original_copyfile(source, destination, *args, **kwargs)

        monkeypatch.setattr(store_module.shutil, "copyfile", fail_metadata_copy)
        with pytest.raises(OSError, match="simulated metadata persistence failure"):
            store.save()

        assert index_path.read_bytes() == previous_index
        assert meta_path.read_bytes() == previous_metadata
        loaded = VectorStore(
            index_path=str(index_path),
            meta_path=str(meta_path),
            embedder=embedder,
        )
        loaded.load()
        assert loaded.index.ntotal == len(loaded.metadata) == 1
        assert loaded.metadata["0"]["document_id"] == "doc-1"

    def test_index_file_restores_memory_when_save_fails(self, tmp_data_dir, monkeypatch):
        class DeterministicEmbedder:
            model_name = "test-embedder"
            dimension = 2

            def encode_chunks(self, chunks):
                return np.array([[float(len(chunk.text)), 0.0] for chunk in chunks], dtype=np.float32)

            def encode_texts(self, texts):
                return np.array([[float(len(text)), 0.0] for text in texts], dtype=np.float32)

            def encode_query(self, query):
                return np.ones(2, dtype=np.float32)

        path = tmp_data_dir / "save-failure.txt"
        _write_txt(path, "Original persisted document content.")
        index_path = str(tmp_data_dir / "save-failure.faiss")
        meta_path = str(tmp_data_dir / "save-failure.json")
        store = VectorStore(
            index_path=index_path,
            meta_path=meta_path,
            embedder=DeterministicEmbedder(),
        )
        index_file(str(path), store, save=True)
        original_metadata = {key: value.copy() for key, value in store.metadata.items()}
        original_vectors = [
            store.index.reconstruct(row).tolist()
            for row in range(store.index.ntotal)
        ]

        _write_txt(path, "Replacement document content which must not commit.")

        def fail_save():
            raise OSError("simulated save failure")

        monkeypatch.setattr(store, "save", fail_save)
        with pytest.raises(OSError, match="simulated save failure"):
            index_file(str(path), store, save=True)

        assert store.metadata == original_metadata
        assert store.index.ntotal == len(original_metadata)
        assert [
            store.index.reconstruct(row).tolist()
            for row in range(store.index.ntotal)
        ] == original_vectors

    def test_failed_replacement_preserves_existing_vectors_and_metadata(self, tmp_data_dir):
        class ToggleEmbedder:
            model_name = "test-embedder"
            dimension = 2
            fail = False

            def encode_chunks(self, chunks):
                if self.fail:
                    raise RuntimeError("simulated embedding failure")
                return np.array([[float(len(chunk.text)), 1.0] for chunk in chunks], dtype=np.float32)

            def encode_texts(self, texts):
                return np.array([[float(len(text)), 1.0] for text in texts], dtype=np.float32)

            def encode_query(self, query):
                return np.array([1.0, 0.0], dtype=np.float32)

        embedder = ToggleEmbedder()
        store = VectorStore(
            index_path=str(tmp_data_dir / "atomic.faiss"),
            meta_path=str(tmp_data_dir / "atomic.json"),
            embedder=embedder,
        )

        def document_chunk(document_id, filename, content):
            source = str(tmp_data_dir / filename)
            document = make_document(
                filename=filename,
                file_type=".txt",
                source=source,
                text=content,
                document_id=document_id,
            )
            chunk = make_chunk(
                document_id=document_id,
                document_name=filename,
                chunk_index=0,
                text=content,
                source_location=source,
            )
            return document, [chunk]

        document_a, chunks_a = document_chunk("doc-a", "a.txt", "original A")
        document_b, chunks_b = document_chunk("doc-b", "b.txt", "original B")
        replacement_a, replacement_chunks_a = document_chunk(
            "doc-a", "a.txt", "replacement A"
        )
        store.index_document(document_a, chunks_a)
        store.index_document(document_b, chunks_b)
        before = {row: record.copy() for row, record in store.metadata.items()}

        embedder.fail = True
        with pytest.raises(RuntimeError, match="simulated embedding failure"):
            store.index_document(replacement_a, replacement_chunks_a)

        assert store.index.ntotal == 2
        assert store.metadata == before
        assert [store.metadata[str(row)]["text"] for row in range(2)] == [
            "original A",
            "original B",
        ]

    def test_reindexing_earlier_source_keeps_rows_aligned(self, tmp_data_dir):
        class DeterministicEmbedder:
            model_name = "test-embedder"
            dimension = 2

            def encode_chunks(self, chunks):
                return self.encode_texts([chunk.text for chunk in chunks])

            def encode_texts(self, texts):
                return np.array(
                    [[float(len(text)), 1.0] for text in texts],
                    dtype=np.float32,
                )

            def encode_query(self, query):
                return np.array([1.0, 0.0], dtype=np.float32)

        store = VectorStore(
            index_path=str(tmp_data_dir / "reindex.faiss"),
            meta_path=str(tmp_data_dir / "reindex.json"),
            embedder=DeterministicEmbedder(),
        )

        def document_chunk(document_id, filename, content):
            source = str(tmp_data_dir / filename)
            document = make_document(
                filename=filename,
                file_type=".txt",
                source=source,
                text=content,
                document_id=document_id,
            )
            chunk = make_chunk(
                document_id=document_id,
                document_name=filename,
                chunk_index=0,
                text=content,
                source_location=source,
            )
            return document, [chunk]

        first_document, first_chunks = document_chunk("doc-a", "a.txt", "first")
        second_document, second_chunks = document_chunk("doc-b", "b.txt", "second")
        replacement_document, replacement_chunks = document_chunk(
            "doc-a", "a.txt", "replacement"
        )

        store.index_document(first_document, first_chunks)
        store.index_document(second_document, second_chunks)
        store.index_document(replacement_document, replacement_chunks)

        assert store.index.ntotal == 2
        assert len(store.metadata) == 2
        assert sorted(store.metadata) == ["0", "1"]
        hits = store.search("query", top_k=2)
        assert {hit["document_id"] for hit in hits} == {"doc-a", "doc-b"}

    def test_index_and_search(self, tmp_data_dir, embedder):
        path = tmp_data_dir / "leave.txt"
        _write_txt(path, "Employees receive twenty days of annual leave per year.")
        index_path = tmp_data_dir / "idx.faiss"
        meta_path = tmp_data_dir / "meta.json"
        store = VectorStore(
            index_path=str(index_path),
            meta_path=str(meta_path),
            embedder=embedder,
        )
        result = index_file(str(path), store, domain="HR", save=True)
        assert result["chunks_added"] >= 1
        store.load()
        hits = store.search("How many leave days?", top_k=1)
        assert hits
        assert hits[0]["filename"] == "leave.txt"
        assert hits[0]["domain"] == "HR"
        assert "embedding_model" in hits[0]

    def test_no_duplicate_indexing_same_source(self, tmp_data_dir, embedder):
        path = tmp_data_dir / "dup.txt"
        _write_txt(path, "Duplicate document content for indexing test.")
        index_path = tmp_data_dir / "dup.faiss"
        meta_path = tmp_data_dir / "dup_meta.json"
        store = VectorStore(
            index_path=str(index_path),
            meta_path=str(meta_path),
            embedder=embedder,
        )
        first = index_file(str(path), store, save=False)
        count_after_first = store.index.ntotal
        second = index_file(str(path), store, save=True)
        assert second["reindexed"] is True
        assert store.index.ntotal == count_after_first
        assert first["chunks_added"] == second["chunks_added"]


class TestValidation:
    def test_supported_extensions(self):
        assert SUPPORTED_EXTENSIONS == frozenset({".pdf", ".docx", ".txt", ".csv"})
