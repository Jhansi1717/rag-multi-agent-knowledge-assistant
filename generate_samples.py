import os
import pandas as pd
from fpdf import FPDF
from docx import Document

os.makedirs('data/hr', exist_ok=True)
os.makedirs('data/technical', exist_ok=True)

# 1. HR PDF: Leave Policy
pdf = FPDF()
pdf.add_page()
pdf.set_font("Arial", size=12)
pdf.cell(200, 10, txt="Company Leave Policy", ln=1, align='C')
pdf.multi_cell(0, 10, txt="""
1. Annual Leave Entitlement
All full-time employees are entitled to 20 days of paid annual leave per calendar year. Part-time employees receive prorated leave based on their working hours.

2. Leave Types
- Annual Leave: For vacation and personal time.
- Sick Leave: 10 days per year for medical reasons.
- Maternity/Paternity Leave: 12 weeks for new parents.
- Bereavement Leave: 3 days in case of the loss of an immediate family member.

3. Leave Application Procedure
Employees must apply for leave through the HR portal at least two weeks in advance for planned absences. Sick leave should be reported to the manager by 9:00 AM on the day of absence.

4. Leave Carry-Forward Rules
Employees may carry forward a maximum of 5 days of unused annual leave to the next calendar year. These carried-forward days must be used by March 31st of the following year.
""")
pdf.output("data/hr/leave_policy.pdf")

# 2. HR CSV: Leave Allowance
data = {
    'Leave_Type': ['Annual', 'Sick', 'Maternity', 'Paternity', 'Bereavement'],
    'Days_Allowed': [20, 10, 84, 84, 3],
    'Carry_Forward_Limit': [5, 0, 0, 0, 0],
    'Approval_Required': ['Manager', 'Manager', 'HR', 'HR', 'Manager']
}
df = pd.DataFrame(data)
df.to_csv("data/hr/leave_allowance.csv", index=False)

# 3. Technical DOCX: API Documentation
doc = Document()
doc.add_heading('API Integration Guidelines', 0)

doc.add_heading('1. Overview', level=1)
doc.add_paragraph('This document outlines the standard approaches for API integration, focusing on REST and SOAP protocols.')

doc.add_heading('2. REST APIs', level=1)
doc.add_paragraph('We primarily use REST (Representational State Transfer) for our microservices architecture. It uses standard HTTP methods for resource manipulation.')

doc.add_heading('3. HTTP Methods', level=1)
doc.add_paragraph('- GET: Retrieve a resource.')
doc.add_paragraph('- POST: Create a new resource.')
doc.add_paragraph('- PUT: Update an existing resource completely.')
doc.add_paragraph('- PATCH: Partially update a resource.')
doc.add_paragraph('- DELETE: Remove a resource.')

doc.add_heading('4. SOAP Services', level=1)
doc.add_paragraph('SOAP (Simple Object Access Protocol) is maintained only for legacy enterprise system integrations. All new services should use REST. SOAP messages are XML-based and typically transported over HTTP.')

doc.add_heading('5. API Authentication', level=1)
doc.add_paragraph('All API endpoints require authentication using OAuth 2.0. Clients must obtain a bearer token and include it in the Authorization header: "Authorization: Bearer <token>".')

doc.save('data/technical/api_documentation.docx')

# 4. Technical TXT: Deployment Guide
txt_content = """
Deployment Procedure for Backend Services

1. Pre-deployment Checklist:
- Ensure all unit and integration tests pass successfully in the CI pipeline.
- Verify that the target environment has the required database schema changes applied.
- Confirm that the Docker image has been built and pushed to the container registry.

2. Deployment Steps:
a. Log in to the deployment server using SSH.
b. Pull the latest Docker image: `docker pull registry.internal/backend-service:latest`
c. Stop the currently running container: `docker stop backend-service`
d. Remove the old container: `docker rm backend-service`
e. Start the new container: 
   `docker run -d --name backend-service -p 8080:8080 --env-file .env registry.internal/backend-service:latest`
f. Check the logs for any startup errors: `docker logs backend-service -f`

3. Post-deployment:
- Run the health check endpoint: `curl -I http://localhost:8080/health`
- Monitor system metrics in the observability dashboard for 15 minutes.
- If errors spike, initiate a rollback to the previous version immediately.
"""
with open("data/technical/deployment_guide.txt", "w") as f:
    f.write(txt_content)

print("Sample documents generated successfully.")
